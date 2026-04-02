"""Generate a 10-slide technical deep-dive deck (Kay style)."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


ROOT = Path("/Users/yilu/Documents/GitHub/idaes-gtep")
DATA_JSON = ROOT / "quality_reports/decks/data/model_update_curtailment_deepdive_data.json"
OUT_PPTX = ROOT / "quality_reports/decks/2026-03-11_Model_Update_and_Curtailment_DeepDive.pptx"

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MEDIUM_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
LIGHT_BLUE = RGBColor(0xD6, 0xE8, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)
AMBER = RGBColor(0xF3, 0x9C, 0x12)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def pct(v: float) -> str:
    return f"{100.0 * v:.2f}%"


def add_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title_text, subtitle_text=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(30)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if subtitle_text:
        shape2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(1.2), SLIDE_W, Inches(0.45)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = MEDIUM_BLUE
        shape2.line.fill.background()
        tf2 = shape2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(15)
        p2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
        p2.font.italic = True
        p2.alignment = PP_ALIGN.LEFT
        tf2.margin_left = Inches(0.6)
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        return Inches(1.85)
    return Inches(1.4)


def add_bullets(slide, left, top, width, height, bullets, font_size=17):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = b.get("level", 0)
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = b["text"]
        run.font.size = Pt(font_size)
        run.font.bold = b.get("bold", False)
        run.font.color.rgb = b.get("color", BLACK)


def add_table(slide, left, top, width, row_h_emu, headers, rows, font_size=11):
    shp = slide.shapes.add_table(
        len(rows) + 1, len(headers), left, top, width, Emu(row_h_emu * (len(rows) + 1))
    )
    tbl = shp.table
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = DARK_BLUE
        p = c.text_frame.paragraphs[0]
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        for j, v in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = bg
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(font_size)
            p.font.color.rgb = BLACK
            p.alignment = PP_ALIGN.CENTER
            c.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_note(slide, txt: str):
    notes = slide.notes_slide.notes_text_frame
    if notes is None:
        return
    notes.clear()
    notes.text = txt


def main():
    payload = json.loads(DATA_JSON.read_text(encoding="utf-8"))
    model = payload["model_update"]
    rows = payload["curtailment_results"]["rows"]
    src = payload["sources"]
    prescient = model["prescient_baseline"]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(
        s,
        "Model Update and Curtailment Analysis Deep Dive",
        "Prescient PCM vs Paper UC | TX-123BT | March 2026",
    )
    add_bullets(
        s,
        Inches(0.8),
        y + Inches(0.25),
        Inches(12.0),
        Inches(4.8),
        [
            {"text": "Goal: update model-comparison interpretation and integrate curtailment-sweep results."},
            {"text": "Inputs: corrected model-comparison report + curtailment notebook exports."},
            {"text": "Output: prioritized benchmark actions and decision points."},
        ],
        21,
    )

    # 2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "What Changed Since Last Update", "Accuracy-review corrections now integrated")
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(0.2),
        Inches(12.2),
        Inches(4.9),
        [
            {"text": "Paper curtailment mechanism corrected: bounded net-load slack, no direct curtailment objective term.", "bold": True},
            {"text": "Solver caveat added: paper shared code uses CONOPT path (continuous relaxation risk)."},
            {"text": "DA-first LMP convention enforced for summarizer (`LMP DA` fallback to `LMP`)."},
            {"text": "Load-weighted LMP metric added for system-level reporting consistency."},
        ],
        17,
    )

    # 3
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Highest-Impact Model Differences", "From current model-comparison ranking")
    top = model["impact_ranking"][:5]
    table_rows = [[r["rank"], r["difference"], r["impact"], r["confidence"]] for r in top]
    add_table(
        s,
        Inches(0.5),
        y + Inches(0.25),
        Inches(12.3),
        420000,
        ["Rank", "Difference", "Impact", "Confidence"],
        table_rows,
        12,
    )
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(3.6),
        Inches(12.1),
        Inches(1.4),
        [{"text": "Structural assumptions remain dominant; cap tuning alone is insufficient.", "bold": True, "color": RED}],
        16,
    )

    # 4
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Paper Curtailment Mechanism (Corrected)", "Code-backed interpretation")
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(0.2),
        Inches(12.2),
        Inches(4.9),
        [
            {"text": "Wind/solar are subtracted from demand first to form net load (`load_b_t`)."},
            {"text": "`rnwcur_b_t` is a nonnegative slack on nodal balance.", "bold": True},
            {"text": "If net load >= 0, `rnwcur_b_t = 0`; if net load < 0, `rnwcur_b_t <= -load_b_t/BaseMVA`."},
            {"text": "Objective excludes direct curtailment cost term."},
            {"text": "Meaning: paper curtailment is balancing relief, not explicit renewable market-bid economics.", "color": AMBER, "bold": True},
        ],
        16,
    )

    # 5
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Our Prescient Curtailment Realization", "Threshold-proxy mechanism")
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(0.2),
        Inches(12.2),
        Inches(4.9),
        [
            {"text": "Current pipeline does not expose a single standalone renewable-curtailment penalty switch."},
            {"text": "Curtailment behavior emerges from SCED/RUC with configured threshold penalties.", "bold": True},
            {"text": "Tracked outputs: `renewables_detail.csv` (`Curtailment`), `hourly_summary.csv` (`RenewablesCurtailment`)."},
            {"text": "Interpretation: this sweep is a diagnostic proxy of economics, not a structural reformulation."},
        ],
        16,
    )

    # 6
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Curtailment Sweep Setup", "Case design and corrected baseline caveat")
    add_table(
        s,
        Inches(0.6),
        y + Inches(0.25),
        Inches(5.0),
        430000,
        ["Penalty ($/MWh)"],
        [[int(r["penalty"])] for r in rows],
        12,
    )
    add_bullets(
        s,
        Inches(6.0),
        y + Inches(0.25),
        Inches(6.9),
        Inches(4.9),
        [
            {"text": "`penalty_1000` is not identical to historical baseline thresholds.", "bold": True, "color": RED},
            {"text": "Uniform-threshold benchmark set is designed for clean sensitivity, not 1:1 baseline reproduction."},
            {"text": "Common comparison window: "
                     f"{payload['curtailment_results']['common_start']} to {payload['curtailment_results']['common_end']}"},
        ],
        15,
    )

    # 7
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Curtailment Results Table", "Core metrics from notebook exports")
    trows = []
    for r in rows:
        trows.append(
            [
                r["case"],
                f"{int(r['penalty'])}",
                pct(r["neg_lmp_frac"]),
                pct(r["floor_hit_frac"]),
                f"{r['weighted_lmp']:.2f}",
                f"{r['total_curtailment_mwh']:.0f}",
                f"{r['total_overgeneration_mwh']:.0f}",
            ]
        )
    add_table(
        s,
        Inches(0.3),
        y + Inches(0.2),
        Inches(12.8),
        350000,
        ["Case", "Penalty", "Neg LMP %", "Floor-hit %", "Weighted LMP", "Curtailment MWh", "OverGen MWh"],
        trows,
        10,
    )

    # 8
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Trend Interpretation", "What the sweep says technically")
    low = rows[0]
    high = rows[-1]
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(0.2),
        Inches(12.2),
        Inches(4.9),
        [
            {
                "text": (
                    f"Neg LMP share increases from {pct(low['neg_lmp_frac'])} "
                    f"(penalty {int(low['penalty'])}) to {pct(high['neg_lmp_frac'])} "
                    f"(penalty {int(high['penalty'])})."
                )
            },
            {"text": "Floor-hit share stays ~1.5% across cases, indicating persistent clipping behavior."},
            {"text": f"Weighted LMP declines from {low['weighted_lmp']:.2f} to {high['weighted_lmp']:.2f} $/MWh.", "color": RED, "bold": True},
            {"text": "Overgeneration remains high and nearly flat (~888-890 GWh), so structural surplus persists."},
        ],
        16,
    )

    # 9
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Benchmark Implications", "What cap tuning can and cannot solve")
    add_bullets(
        s,
        Inches(0.6),
        y + Inches(0.2),
        Inches(12.2),
        Inches(4.9),
        [
            {"text": "Cap tuning is a useful diagnostic, not a full benchmark-alignment solution."},
            {"text": "Persistent negative-price behavior points to formulation mismatch, not only threshold settings.", "bold": True},
            {"text": "Primary levers remain: renewable representation, chronology coupling, and pricing workflow."},
        ],
        16,
    )

    # 10
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(s)
    y = add_title_bar(s, "Action Plan and Decision Points", "Next execution sequence")
    add_bullets(
        s,
        Inches(0.7),
        y + Inches(0.25),
        Inches(12.0),
        Inches(4.9),
        [
            {"text": "1. Approve paper-like benchmark mode for renewables/net-load treatment.", "bold": True},
            {"text": "2. Keep curtailment sweep as diagnostics; do not treat it as structural equivalence."},
            {"text": "3. Run day-isolated benchmark window (e.g., Day110) for sign/trend alignment check."},
            {"text": "4. Define acceptance criteria: sign/trend first, magnitude second."},
        ],
        17,
    )

    for slide in prs.slides:
        add_note(
            slide,
            "Sources: "
            f"{src['model_comparison_report']}; "
            f"{src['curtailment_setup_report']}; "
            f"{src['curtailment_summary_csv']}; "
            f"{src['curtailment_notebook']}",
        )

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    print(f"Saved: {OUT_PPTX}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
