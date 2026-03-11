"""Generate benchmark update presentation using Kay deck visual style."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


ROOT = Path("/Users/yilu/Documents/GitHub/idaes-gtep")
DATA_JSON = ROOT / "quality_reports/decks/data/benchmark_update_slide_data.json"
OUT_PPTX = ROOT / "quality_reports/decks/2026-03-03_Benchmarking_Update_PCM_vs_Paper.pptx"

# Colors (aligned with create_presentation.py)
DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MEDIUM_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
LIGHT_BLUE = RGBColor(0xD6, 0xE8, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x27, 0xAE, 0x60)
AMBER = RGBColor(0xF3, 0x9C, 0x12)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title_text, subtitle_text=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(31)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    if subtitle_text:
        shape2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(1.2), SLIDE_W, Inches(0.45)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = MEDIUM_BLUE
        shape2.line.fill.background()
        tf2 = shape2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
        p2.font.italic = True
        p2.alignment = PP_ALIGN.LEFT
        tf2.margin_left = Inches(0.6)
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        return Inches(1.85)
    return Inches(1.4)


def add_bullet_box(slide, left, top, width, height, bullets, font_size=18):
    tx_box = slide.shapes.add_textbox(left, top, width, height)
    tf = tx_box.text_frame
    tf.word_wrap = True
    for idx, b in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.level = b.get("level", 0)
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = b["text"]
        run.font.size = Pt(font_size)
        run.font.bold = b.get("bold", False)
        run.font.color.rgb = b.get("color", BLACK)
    return tx_box


def add_table(slide, left, top, width, row_height_emu, headers, rows):
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), left, top, width, Emu(row_height_emu * (len(rows) + 1))
    )
    table = shape.table

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    for ridx, row in enumerate(rows):
        bg = LIGHT_BLUE if ridx % 2 == 0 else WHITE
        for cidx, val in enumerate(row):
            cell = table.cell(ridx + 1, cidx)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(11)
            p.font.color.rgb = BLACK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape


def add_note(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def add_card(slide, left, top, width, height, title, body, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = WHITE
    p1.alignment = PP_ALIGN.LEFT
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(13)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.LEFT


def pct(v):
    return f"{v * 100:.2f}%"


def main():
    payload = json.loads(DATA_JSON.read_text(encoding="utf-8"))
    model = payload["model_comparison"]
    curtail = payload["curtailment"]
    hydro = payload["hydro_q1"]["scenarios"]
    sources = payload["sources"]

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # Slide 1
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(
        slide,
        "Prescient PCM vs Paper UC: Benchmarking Status and Effort",
        "TX-123BT ERCOT | Collaborator/Supervisor Update | March 2026",
    )
    add_bullet_box(
        slide,
        Inches(0.8),
        y + Inches(0.3),
        Inches(12.0),
        Inches(3.8),
        [
            {"text": "Purpose: estimate effort needed to align Prescient PCM with paper UC benchmark."},
            {"text": "Coverage: model-difference audit, curtailment experiment status, hydro Q1 benchmark signal."},
            {"text": "Decision focus: where to invest benchmarking effort next."},
        ],
        font_size=22,
    )
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(6.85), SLIDE_W, Inches(0.65))
    footer.fill.solid()
    footer.fill.fore_color.rgb = DARK_BLUE
    footer.line.fill.background()
    fp = footer.text_frame.paragraphs[0]
    fp.text = "IDAES-GTEP | Benchmarking update section (~20 minutes)"
    fp.font.size = Pt(15)
    fp.font.color.rgb = WHITE
    fp.alignment = PP_ALIGN.CENTER
    add_note(slide, f"Sources: {sources['model_comparison_report']}; {sources['curtailment_setup_report']}")

    # Slide 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Executive Findings", "Structural mismatch dominates current LMP divergence")
    add_bullet_box(
        slide,
        Inches(0.5),
        y + Inches(0.2),
        Inches(6.0),
        Inches(4.8),
        [
            {"text": "Paper UC sampled outputs are all-positive LMP."},
            {
                "text": (
                    "Prescient baseline still shows negative-LMP share "
                    f"~{pct(model['prescient_baseline']['neg_frac'])} with -1000 floor hits."
                ),
                "color": RED,
            },
            {"text": "Top drivers are formulation differences: renewable treatment, chronology, pricing workflow."},
            {"text": "Hydro inclusion in Q1 did not materially improve alignment to paper behavior.", "color": AMBER},
        ],
        font_size=17,
    )
    add_card(
        slide,
        Inches(6.8),
        y + Inches(0.2),
        Inches(2.0),
        Inches(2.1),
        "Finding 1",
        "Mismatch is systematic,\nnot random noise.",
        DARK_BLUE,
    )
    add_card(
        slide,
        Inches(9.0),
        y + Inches(0.2),
        Inches(2.0),
        Inches(2.1),
        "Finding 2",
        "Hydro is not the\nprimary lever now.",
        MEDIUM_BLUE,
    )
    add_card(
        slide,
        Inches(11.2),
        y + Inches(0.2),
        Inches(1.9),
        Inches(2.1),
        "Finding 3",
        "Need paper-like\nbenchmark mode.",
        GRAY,
    )
    add_note(slide, f"Sources: {sources['model_comparison_report']}; {sources['hydro_q1_summary_json']}")

    # Slide 3
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Paper UC vs Prescient PCM: Comparison Matrix", "Why price behavior diverges")
    rows = [
        ["Renewables", "Subtracted from load (net-load)", "Explicit generators in clearing", "Very high"],
        ["Curtailment", "Free relief variable in balance", "Penalty/cap mediated behavior", "Very high"],
        ["Pricing", "Two-pass dual pricing", "Integrated RUC/SCED pricing", "High"],
        ["Chronology", "Daily independent runs", "Rolling horizons (90d, look-ahead)", "High"],
        ["Reserves", "Custom reserve constraints", "reserve_factor design", "Medium"],
        ["Line ratings", "Day-specific / hourly DLR options", "RTS-GMLC path unless remapped", "Medium"],
    ]
    add_table(
        slide,
        Inches(0.5),
        y + Inches(0.2),
        Inches(12.3),
        340000,
        ["Aspect", "Paper model", "Our Prescient", "LMP impact"],
        rows,
    )
    add_note(slide, f"Source: {sources['model_comparison_report']}")

    # Slide 4
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Evidence Snapshot: LMP Divergence", "Sign behavior differs structurally")
    add_table(
        slide,
        Inches(0.7),
        y + Inches(0.35),
        Inches(12.0),
        420000,
        ["Case", "Min", "Max", "Mean", "Negative LMP share"],
        [
            ["Paper UC (sampled days)", ">=0", "45-46", "positive", "0.00%"],
            [
                "Prescient baseline",
                f"{model['prescient_baseline']['min']:.1f}",
                f"{model['prescient_baseline']['max']:.1f}",
                f"{model['prescient_baseline']['mean']:.2f}",
                pct(model["prescient_baseline"]["neg_frac"]),
            ],
            [
                "Prescient UC-only",
                f"{model['prescient_uc_only']['min']:.1f}",
                f"{model['prescient_uc_only']['max']:.1f}",
                f"{model['prescient_uc_only']['mean']:.2f}",
                pct(model["prescient_uc_only"]["neg_frac"]),
            ],
        ],
    )
    add_bullet_box(
        slide,
        Inches(0.7),
        y + Inches(3.3),
        Inches(12.0),
        Inches(1.8),
        [
            {"text": "Observed -1000 floor and ~11% negative share indicate design-assumption mismatch, not sampling noise."},
            {"text": "Priority should be formulation alignment before tuning scenario details.", "bold": True, "color": DARK_BLUE},
        ],
        font_size=16,
    )
    add_note(slide, f"Source: {sources['model_comparison_report']}")

    # Slide 5
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Curtailment Penalty Work: Current Status", "Experiment framework in place; run still in progress")
    add_bullet_box(
        slide,
        Inches(0.5),
        y + Inches(0.2),
        Inches(7.7),
        Inches(4.8),
        [
            {"text": curtail["status"], "bold": True, "color": AMBER},
            {"text": "Experiment implemented in current Prescient pipeline and reproducible by case manifest."},
            {"text": "Purpose: measure sensitivity of curtailment, negative-LMP share, and floor-hit behavior."},
            {"text": f"Runner script: {curtail['script']}"},
            {"text": "This section intentionally remains brief until the full run is complete."},
        ],
        font_size=16,
    )
    add_table(
        slide,
        Inches(8.4),
        y + Inches(0.4),
        Inches(4.4),
        420000,
        ["Penalty sweep ($/MWh)"],
        [[v] for v in curtail["penalty_values_usd_per_mwh"]],
    )
    add_note(slide, f"Source: {sources['curtailment_setup_report']}")

    # Slide 6
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Hydro Q1 Check: No Material Improvement", "Hydro inclusion does not uplift LMP alignment in current setup")
    add_table(
        slide,
        Inches(0.5),
        y + Inches(0.3),
        Inches(12.3),
        390000,
        ["Scenario", "Weighted LMP", "Negative LMP share", "Floor-hit share", "Min", "Max"],
        [
            [
                "Hydro Btheta UC+ED",
                f"{hydro['Hydro Btheta UC+ED']['lmp_weighted']:.2f}",
                pct(hydro["Hydro Btheta UC+ED"]["neg_lmp_frac"]),
                pct(hydro["Hydro Btheta UC+ED"]["floor_hit_frac"]),
                f"{hydro['Hydro Btheta UC+ED']['lmp_min']:.1f}",
                f"{hydro['Hydro Btheta UC+ED']['lmp_max']:.1f}",
            ],
            [
                "Hydro Btheta UC-only",
                f"{hydro['Hydro Btheta UC-only']['lmp_weighted']:.2f}",
                pct(hydro["Hydro Btheta UC-only"]["neg_lmp_frac"]),
                pct(hydro["Hydro Btheta UC-only"]["floor_hit_frac"]),
                f"{hydro['Hydro Btheta UC-only']['lmp_min']:.1f}",
                f"{hydro['Hydro Btheta UC-only']['lmp_max']:.1f}",
            ],
            [
                "No-hydro Benchmark Q1",
                f"{hydro['No-hydro Benchmark Q1']['lmp_weighted']:.2f}",
                pct(hydro["No-hydro Benchmark Q1"]["neg_lmp_frac"]),
                pct(hydro["No-hydro Benchmark Q1"]["floor_hit_frac"]),
                f"{hydro['No-hydro Benchmark Q1']['lmp_min']:.1f}",
                f"{hydro['No-hydro Benchmark Q1']['lmp_max']:.1f}",
            ],
        ],
    )
    add_bullet_box(
        slide,
        Inches(0.6),
        y + Inches(3.35),
        Inches(12.0),
        Inches(1.6),
        [
            {"text": payload["hydro_q1"]["summary"], "bold": True, "color": RED},
            {"text": "Hydro addition should not be treated as the primary benchmark-alignment lever at this stage."},
        ],
        font_size=15,
    )
    add_note(slide, f"Sources: {sources['hydro_q1_summary_json']}; {sources['hydro_q1_notebook']}")

    # Slide 7
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Effort-to-Alignment Roadmap", "Estimated 4-6 weeks for high-confidence benchmark alignment")
    add_table(
        slide,
        Inches(0.5),
        y + Inches(0.3),
        Inches(12.3),
        340000,
        ["Phase", "Change", "Deliverable", "Est. effort", "Risk"],
        [
            ["1", "Paper-like benchmark mode", "Day110 sign/trend match check", "1-1.5 weeks", "Medium"],
            ["2", "Pricing + cap sensitivity", "Floor-hit reduction diagnostics", "1 week", "Low-Medium"],
            ["3", "Reserve + line-rating alignment", "Trend-shape stability check", "1-1.5 weeks", "Medium"],
            ["4", "Quarter/full-year robustness", "Benchmark report and confidence band", "1-2 weeks", "Medium-High"],
        ],
    )
    add_bullet_box(
        slide,
        Inches(0.6),
        y + Inches(3.2),
        Inches(12.0),
        Inches(1.8),
        [
            {"text": "Checkpoint rule: proceed phase-by-phase only if negative-LMP share and trend gaps improve materially."},
            {"text": "Expected total effort: 4-6 focused weeks.", "bold": True, "color": DARK_BLUE},
        ],
        font_size=16,
    )
    add_note(slide, "Roadmap synthesized from model comparison report and hydro/curtailment status.")

    # Slide 8
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(slide, "Decisions Requested + Next 2 Weeks", "Lock scope before deeper simulation spend")
    add_bullet_box(
        slide,
        Inches(0.7),
        y + Inches(0.25),
        Inches(12.0),
        Inches(5.0),
        [
            {"text": "Decision 1: Approve dedicated paper-like benchmark mode branch.", "bold": True},
            {"text": "Decision 2: Prioritize structural alignment (renewables/chronology/pricing) before hydro expansion.", "bold": True},
            {"text": "Decision 3: Define success criteria as sign/trend alignment first, magnitude second.", "bold": True},
            {"text": "Week 1 deliverable: Day110 benchmark package with side-by-side LMP diagnostics."},
            {"text": "Week 2 deliverable: cap-sensitivity and reserve/rating alignment impact memo."},
        ],
        font_size=17,
    )
    add_note(slide, "This closing slide is for meeting decisions and near-term execution commitments.")

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    print(f"Saved: {OUT_PPTX}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
