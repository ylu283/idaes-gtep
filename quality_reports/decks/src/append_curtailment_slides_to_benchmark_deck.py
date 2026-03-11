"""Append curtailment update slides to the existing benchmark deck.

Append-only update for:
quality_reports/decks/2026-03-03_Benchmarking_Update_PCM_vs_Paper.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt


ROOT = Path("/Users/yilu/Documents/GitHub/idaes-gtep")
DECK_PATH = ROOT / "quality_reports/decks/2026-03-03_Benchmarking_Update_PCM_vs_Paper.pptx"

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MEDIUM_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
LIGHT_BLUE = RGBColor(0xD6, 0xE8, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
RED = RGBColor(0xC0, 0x39, 0x2B)

SLIDE_W = Inches(13.333)


def add_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title_text, subtitle_text=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(30)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    if subtitle_text:
        shape2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(1.2), SLIDE_W, Inches(0.45)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = MEDIUM_BLUE
        shape2.line.fill.background()
        tf2 = shape2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(15)
        p2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
        p2.font.italic = True
        p2.alignment = PP_ALIGN.LEFT
        tf2.margin_left = Inches(0.6)
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        return Inches(1.85)
    return Inches(1.4)


def add_bullets(slide, left, top, width, height, lines, font_size=17):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = item.get("level", 0)
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = item["text"]
        run.font.size = Pt(font_size)
        run.font.bold = item.get("bold", False)
        run.font.color.rgb = item.get("color", BLACK)


def add_table(slide, left, top, width, row_h_emu, headers, rows, font_size=10):
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), left, top, width, Emu(row_h_emu * (len(rows) + 1))
    )
    tbl = shape.table
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = DARK_BLUE
        p = c.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        for j, val in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.text = str(val)
            c.fill.solid()
            c.fill.fore_color.rgb = bg
            p = c.text_frame.paragraphs[0]
            p.font.size = Pt(font_size)
            p.font.color.rgb = BLACK
            p.alignment = PP_ALIGN.CENTER
            c.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_note(slide, text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    if tf is None:
        return
    tf.clear()
    tf.text = text


def main():
    prs = Presentation(str(DECK_PATH))

    # Slide 13
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(
        slide,
        "How Curtailment Is Realized in Our Prescient Pipeline",
        "Proxy mechanism via SCED/RUC penalty-threshold economics",
    )
    add_bullets(
        slide,
        Inches(0.7),
        y + Inches(0.2),
        Inches(12.0),
        Inches(4.8),
        [
            {"text": "No standalone 'renewable curtailment penalty' switch in current Prescient run configuration.", "bold": True},
            {"text": "We tune price/violation thresholds to shape dispatch and curtailment outcomes."},
            {"text": "Curtailment is tracked directly in outputs:", "bold": True},
            {"text": "renewables_detail.csv -> Curtailment", "level": 1},
            {"text": "hourly_summary.csv -> RenewablesCurtailment", "level": 1},
            {"text": "Sweep cases: 300 / 1000 / 2000 / 5000 / 10000 $/MWh.", "bold": True},
            {"text": "Interpretation caution: this is a curtailment-penalty proxy, not a one-parameter physical curtailment model."},
        ],
        font_size=16,
    )
    add_note(
        slide,
        "Sources: quality_reports/reports/curtailment_penalty_experiment_setup_2026-02-28.md; "
        "gtep/data/retirement_allowed_no_extreme_half_load/run_curtailment_penalty_experiments.py",
    )

    # Slide 14
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(
        slide,
        "Curtailment Penalty Sweep Results (Jan-Mar Common Window)",
        "Common window: 2019-01-01 00:00 to 2019-03-31 23:00",
    )
    headers = [
        "Case",
        "Neg LMP %",
        "Floor-hit %",
        "Weighted LMP",
        "Curtail (MWh)",
        "OverGen (MWh)",
    ]
    rows = [
        ["penalty_300", "9.62%", "1.52%", "9.79", "118,560", "888,328"],
        ["penalty_1000", "10.51%", "1.50%", "3.96", "124,325", "889,003"],
        ["penalty_2000", "11.43%", "1.51%", "-7.91", "125,040", "888,986"],
        ["penalty_5000", "11.91%", "1.49%", "-39.57", "121,351", "889,607"],
        ["penalty_10000", "12.80%", "1.49%", "-88.35", "123,328", "888,693"],
    ]
    add_table(slide, Inches(0.5), y + Inches(0.25), Inches(12.3), 320000, headers, rows, font_size=11)
    add_bullets(
        slide,
        Inches(0.6),
        y + Inches(2.9),
        Inches(12.1),
        Inches(2.2),
        [
            {"text": "Higher caps widen tails and increase negative-LMP frequency.", "color": RED, "bold": True},
            {"text": "Floor-hit fraction stays ~1.5% across cases (persistent clipping)."},
            {"text": "Weighted LMP trends downward sharply as cap increases."},
            {"text": "Overgeneration remains high and nearly flat (~888-890 GWh): cap tuning alone is not a structural fix."},
        ],
        font_size=15,
    )
    add_note(
        slide,
        "Sources: gtep/pcm_analysis/curtailment_penalty_benchmark_summary.csv and "
        "gtep/pcm_analysis/prescient_lmp_analysis_curtailment_penalty.ipynb",
    )

    # Slide 15
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    y = add_title_bar(
        slide,
        "Corrected Paper Curtailment Mechanism vs Our Model",
        "Paper: bounded net-load slack; Ours: threshold-driven market proxy behavior",
    )
    add_bullets(
        slide,
        Inches(0.6),
        y + Inches(0.2),
        Inches(5.9),
        Inches(4.9),
        [
            {"text": "Paper Pyomo (actual implementation):", "bold": True},
            {"text": "Wind/solar subtracted from load first (net-load).", "level": 1},
            {"text": "rnwcur_b_t >= 0 and only active when net load < 0.", "level": 1},
            {"text": "Bounded by magnitude of negative net load.", "level": 1},
            {"text": "Appears in nodal balance RHS as balancing relief.", "level": 1},
            {"text": "No direct curtailment objective-cost term.", "level": 1, "bold": True},
        ],
        font_size=14,
    )
    add_bullets(
        slide,
        Inches(6.8),
        y + Inches(0.2),
        Inches(6.0),
        Inches(4.9),
        [
            {"text": "Our Prescient benchmark run:", "bold": True},
            {"text": "Explicit renewable participation in clearing.", "level": 1},
            {"text": "Curtailment influenced by penalty/threshold economics.", "level": 1},
            {"text": "Cap tuning changes tail behavior strongly.", "level": 1},
            {"text": "But surplus/overgeneration remains structurally high.", "level": 1},
            {"text": "Implication: structural alignment still needed for paper-like LMP behavior.", "level": 1, "bold": True, "color": RED},
        ],
        font_size=14,
    )
    add_note(
        slide,
        "Sources: Sample_Codes_SCUC/UC_function.py, Run_SCUC_annual.py, "
        "Sample_Codes_SCUC_HourlyDLR/UC_function_DLR.py, RunUC_annual_dlr.py; "
        "quality_reports/reports/model_comparison_prescient_vs_paper_uc_2026-02-28.md",
    )

    prs.save(str(DECK_PATH))
    print(f"Appended 3 slides to: {DECK_PATH}")
    print(f"Total slides now: {len(prs.slides)}")


if __name__ == "__main__":
    main()
