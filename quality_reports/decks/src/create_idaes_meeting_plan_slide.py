"""Generate a standalone one-slide IDAES meeting plan deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/Users/yilu/Documents/GitHub/idaes-gtep")
OUT_PPTX = ROOT / "quality_reports/decks/2026-03-26_idaes_meeting_plan_slide.pptx"

DARK_BLUE = RGBColor(0x1B, 0x3A, 0x5C)
MEDIUM_BLUE = RGBColor(0x2C, 0x5F, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_title_bar(slide, title_text, subtitle_text=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()

    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(31)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.6)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    if subtitle_text:
        shape2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(1.2), SLIDE_W, Inches(0.45)
        )
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = MEDIUM_BLUE
        shape2.line.fill.background()

        tf2 = shape2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xEE)
        p2.font.italic = True
        p2.alignment = PP_ALIGN.LEFT
        tf2.margin_left = Inches(0.6)
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE

        return Inches(1.85)

    return Inches(1.4)


def add_bullets(slide, left, top, width, height, lines):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)

        run = p.add_run()
        run.text = line["text"]
        run.font.size = Pt(20)
        run.font.bold = line.get("bold", False)
        run.font.color.rgb = line.get("color", BLACK)


def add_footer(slide, text):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(6.72), Inches(12.2), Inches(0.45)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF2, 0xF6, 0xFA)
    shape.line.color.rgb = RGBColor(0xD9, 0xE2, 0xEC)

    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(13)
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    y = add_title_bar(
        slide,
        "IDAES PCM Benchmarking Plan (Mar 29 - Apr 23, 2026)",
        "Goal: align GTEP-to-PCM conversion and quantify LMP deviation sources vs paper UC",
    )

    bullets = [
        {
            "text": "By Mar 29 (Sun): Compare how GTEP expansion outputs convert to PCM-ready inputs; fact-check mapping logic.",
            "bold": True,
        },
        {
            "text": "Mar 30 - Apr 12 (2 weeks): Run PCM simulations for baseline and converted scenarios with clean manifests/log checks.",
        },
        {
            "text": "Apr 13 - Apr 19 (1 week): Compare results and analyze time-series aggregation error impacts on LMP trend/sign.",
        },
        {
            "text": "Apr 23 (Thu): Present conversion status, PCM benchmark results, and next-step recommendations at bi-weekly IDAES meeting.",
            "bold": True,
        },
    ]

    add_bullets(slide, Inches(0.7), y + Inches(0.28), Inches(12.0), Inches(4.9), bullets)
    add_footer(
        slide,
        "Fact check: Mar 29, 2026 is Sunday; meeting target is Thursday, Apr 23, 2026.",
    )

    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    notes.text = (
        "Prepared as standalone planning slide for IDAES bi-weekly update. "
        "Timeline reflects requested 2-week PCM run and 1-week analysis window."
    )

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"Wrote: {OUT_PPTX}")


if __name__ == "__main__":
    main()
